"""
Aktualisiert taeglich die Anzahl "Unfallfreie Tage" auf der Kiosk-Folie.

Berechnet die Tage seit `startDate` (config.json) neu, statt stur +1 zu
zaehlen. Das ist robuster: verpasste Laeufe (Wochenende, PC aus) korrigieren
sich von selbst, und nach einem neuen Unfall muss nur `startDate` in
config.json angepasst werden.

Einrichtung (einmalig):
  pip install python-pptx

Taeglicher Aufruf (z.B. ueber Windows-Aufgabenplanung oder cron), z.B. um 00:05 Uhr:
  python update_counter.py
"""

import json
from datetime import date, datetime
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.json"
PPTX_PATH = HERE / "Unfallfreie_Tage_DC_Losheim.pptx"


def load_start_date() -> date:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    return datetime.strptime(config["startDate"], "%Y-%m-%d").date()


def update_slide() -> int:
    start_date = load_start_date()
    today = date.today()
    day_count = (today - start_date).days

    prs = Presentation(PPTX_PATH)
    slide = prs.slides[0]

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()

        # The big counter: a text box whose entire content is just digits.
        if text.isdigit():
            shape.text_frame.paragraphs[0].runs[0].text = str(day_count)

        # The footer line with the two dates.
        elif text.startswith("Letzter Unfall:"):
            run = shape.text_frame.paragraphs[0].runs[0]
            run.text = (
                f"Letzter Unfall: {start_date.strftime('%d.%m.%Y')}"
                f"      |      Stand: {today.strftime('%d.%m.%Y')}"
            )

    prs.save(PPTX_PATH)
    return day_count


if __name__ == "__main__":
    count = update_slide()
    print(f"Aktualisiert: {PPTX_PATH.name} -> {count} unfallfreie Tage ({date.today().isoformat()})")
